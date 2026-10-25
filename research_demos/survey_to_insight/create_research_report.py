
"""Create a research report from Survey-to-Insight analysis outputs.

This report uses synthetic demonstration data. It is not evidence about
actual educators, schools, or education institutions.
"""

import argparse
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def add_dataframe_table(document, dataframe, max_rows=20):
    """Add a readable table containing a DataFrame's first rows."""
    if dataframe.empty:
        document.add_paragraph("No records available for this table.")
        return

    display = dataframe.head(max_rows)
    table = document.add_table(rows=1, cols=len(display.columns))
    table.style = "Light Shading Accent 1"

    for cell, heading in zip(table.rows[0].cells, display.columns):
        cell.text = str(heading)

    for _, record in display.iterrows():
        cells = table.add_row().cells
        for cell, value in zip(cells, record):
            if pd.isna(value):
                cell.text = ""
            elif isinstance(value, float):
                cell.text = f"{value:.2f}"
            else:
                cell.text = str(value)

    if len(dataframe) > max_rows:
        document.add_paragraph(
            f"Showing {max_rows} of {len(dataframe)} rows. "
            "See the accompanying CSV outputs for the complete table."
        )


def add_chart_if_available(document, path, caption):
    """Insert a chart if its file exists."""
    if path.exists():
        document.add_picture(str(path), width=Inches(6.0))
        paragraph = document.add_paragraph(caption)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        document.add_paragraph(
            f"Chart not available: {path.name}"
        )


def create_report(analysis_dir: Path, output_file: Path) -> Path:
    """Build a Word report from existing analysis outputs."""

    manifest_path = analysis_dir / "analysis_manifest.json"
    validation_path = analysis_dir / "validation_report.json"

    if not manifest_path.exists():
        raise FileNotFoundError(
            f"Analysis manifest not found: {manifest_path}"
        )
    if not validation_path.exists():
        raise FileNotFoundError(
            f"Validation report not found: {validation_path}"
        )

    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    validation = json.loads(
        validation_path.read_text(encoding="utf-8")
    )

    if manifest.get("dataset_type") != "SYNTHETIC_NOT_REAL":
        raise ValueError(
            "Report generation requires data labelled SYNTHETIC_NOT_REAL."
        )
    if not validation.get("analysis_allowed", False):
        raise ValueError(
            "Report generation stopped because validation does not allow analysis."
        )

    required_tables = {
        "categorical_frequencies.csv": "Sample profile",
        "familiarity_descriptive_statistics.csv": "Skill familiarity statistics",
        "familiarity_means_by_role.csv": "Familiarity means by role",
        "q10_task_frequencies.csv": "Reported task frequencies",
        "q12_response_completeness.csv": "Optional response completeness",
    }

    tables = {}
    for filename, label in required_tables.items():
        path = analysis_dir / filename
        if not path.exists():
            raise FileNotFoundError(
                f"Required analysis table not found: {path}"
            )
        tables[label] = pd.read_csv(path)

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)

    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(10)

    title = document.add_heading(
        "Research and Data Skills Needs Assessment", level=0
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = document.add_paragraph(
        "Synthetic demonstration report — NOT REAL RESPONDENT DATA"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph(
        "Purpose: demonstrate a reproducible workflow for validating "
        "survey data, producing descriptive statistics, and preparing "
        "a decision-oriented research report."
    )

    document.add_heading("1. Executive Summary", level=1)
    document.add_paragraph(
        f"This demonstration analyses {manifest['records_analysed']} "
        "synthetic records. The dataset was generated for software and "
        "analytical workflow demonstration; it was not collected from "
        "actual educators or institutions."
    )
    document.add_paragraph(
        f"Automated validation allowed analysis, with "
        f"{manifest['critical_validation_errors']} critical validation "
        f"errors and {manifest['validation_warnings']} warnings."
    )
    document.add_paragraph(
        "The tables and charts below describe the generated records only. "
        "They must not be interpreted as estimates of real educators' "
        "skills, needs, or experiences."
    )

    document.add_heading("2. Objectives", level=1)
    for objective in [
        "Demonstrate automated survey-data quality checks.",
        "Summarise categorical responses and self-reported familiarity ratings.",
        "Illustrate role-based descriptive comparisons and multiple-response analysis.",
        "Produce traceable outputs that can be reproduced and independently checked.",
    ]:
        document.add_paragraph(objective, style="List Bullet")

    document.add_heading("3. Methods", level=1)
    document.add_paragraph(
        "The example questionnaire covers professional role, experience, "
        "familiarity with research and data skills, activities undertaken "
        "during the previous year, development needs, and an optional "
        "open-text challenge."
    )
    document.add_paragraph(
        "The dataset was generated programmatically. Automated tests "
        "check the data generator, validation rules, analysis outputs, "
        "and workbook safeguards. The analysis uses descriptive "
        "frequencies, percentages, summary statistics, group means, "
        "and charts."
    )
    document.add_paragraph(
        "The questionnaire has not been administered to real respondents "
        "for this demonstration. Its items should not be treated as a "
        "validated psychometric scale."
    )

    document.add_heading("4. Data Quality and Validation", level=1)
    quality_rows = [
        ("Records", validation.get("dataset_rows")),
        ("Columns", validation.get("dataset_columns")),
        ("Analysis allowed", validation.get("analysis_allowed")),
        ("Critical errors", validation.get("critical_error_count")),
        ("Warnings", validation.get("warning_count")),
    ]
    quality_table = document.add_table(rows=1, cols=2)
    quality_table.style = "Light Shading Accent 1"
    quality_table.rows[0].cells[0].text = "Check"
    quality_table.rows[0].cells[1].text = "Result"
    for check, result in quality_rows:
        cells = quality_table.add_row().cells
        cells[0].text = str(check)
        cells[1].text = str(result)

    document.add_paragraph(
        "Validation indicates whether the generated file satisfies the "
        "implemented rules. It does not establish the validity of the "
        "questionnaire or the representativeness of the data."
    )

    document.add_heading("5. Descriptive Results", level=1)

    document.add_heading("5.1 Professional profile and development needs", level=2)
    add_dataframe_table(document, tables["Sample profile"])

    document.add_heading("5.2 Familiarity with research and data skills", level=2)
    add_dataframe_table(
        document, tables["Skill familiarity statistics"]
    )
    add_chart_if_available(
        document,
        analysis_dir / "mean_familiarity_by_skill.png",
        "Figure 1. Mean familiarity ratings in the synthetic dataset.",
    )

    document.add_heading("5.3 Familiarity by professional role", level=2)
    add_dataframe_table(document, tables["Familiarity means by role"])

    document.add_heading("5.4 Activities reported in the generated records", level=2)
    add_dataframe_table(document, tables["Reported task frequencies"])

    document.add_heading("5.5 Optional open-text response completeness", level=2)
    add_dataframe_table(
        document, tables["Optional response completeness"]
    )

    add_chart_if_available(
        document,
        analysis_dir / "respondents_by_role.png",
        "Figure 2. Number of synthetic records by professional role.",
    )

    document.add_heading("6. Interpretation", level=1)
    document.add_paragraph(
        "The descriptive outputs demonstrate how a researcher can "
        "summarise a survey dataset and compare response patterns across "
        "groups. Because the records were simulated, any apparent "
        "differences between roles or skill areas reflect the data "
        "generation process rather than measured differences in a real "
        "population."
    )
    document.add_paragraph(
        "Percentages for the multiple-response task question use the "
        "total number of records as the denominator. Because respondents "
        "can select more than one task, these percentages need not sum "
        "to 100 percent."
    )

    document.add_heading("7. Limitations", level=1)
    for limitation in [
        "All records are synthetic and are not evidence about real educators.",
        "The questionnaire has not been administered or psychometrically validated in this demonstration.",
        "Descriptive associations do not establish causation.",
        "Generated response patterns may not resemble those in a real population.",
        "The optional open-text responses are not a substitute for qualitative analysis of authentic participant accounts.",
    ]:
        document.add_paragraph(limitation, style="List Bullet")

    document.add_heading("8. Recommendations for a Real Assessment", level=1)
    for recommendation in [
        "Review the questionnaire with relevant education and research stakeholders before administration.",
        "Pilot the instrument and refine unclear or overlapping response options.",
        "Define the target population, sampling strategy, consent process, and data-protection procedures.",
        "Collect authorised real data and retain a clear record of its provenance.",
        "Repeat validation before analysis and document missing data and deviations.",
        "Interpret findings in light of sampling limitations and triangulate them with appropriate qualitative evidence.",
        "Use the findings to prioritise training only after the real data and contextual evidence have been reviewed.",
    ]:
        document.add_paragraph(recommendation, style="List Number")

    document.add_heading("9. Reproducibility and Provenance", level=1)
    document.add_paragraph(
        f"Input file: {manifest.get('input_file', 'Not recorded')}"
    )
    document.add_paragraph(
        f"Dataset classification: {manifest.get('dataset_type')}"
    )
    document.add_paragraph(
        "The associated CSV files, JSON validation report, analysis "
        "manifest, charts, and Excel workbook provide supporting outputs. "
        "This report should be read alongside those files."
    )

    output_file.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_file)
    return output_file


def main():
    parser = argparse.ArgumentParser(
        description="Create a Word report from validated survey analysis."
    )
    parser.add_argument(
        "--analysis-dir",
        type=Path,
        default=Path(__file__).resolve().parent / "analysis_outputs",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).resolve().parent
        / "analysis_outputs"
        / "survey_to_insight_research_report.docx",
    )
    args = parser.parse_args()

    output = create_report(args.analysis_dir, args.output)
    print("Research report created successfully.")
    print("Dataset status: SYNTHETIC_NOT_REAL")
    print(f"Report: {output.resolve()}")


if __name__ == "__main__":
    main()
