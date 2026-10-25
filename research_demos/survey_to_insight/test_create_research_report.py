
"""Tests for the Survey-to-Insight Word research report generator."""

import json
from pathlib import Path

import pytest
from docx import Document

from analyse_survey import analyse_survey
from create_research_report import create_report
from generate_synthetic_data import generate_data


def prepare_analysis(tmp_path: Path) -> Path:
    """Generate and analyse temporary synthetic survey data."""
    input_csv = tmp_path / "responses.csv"
    analysis_dir = tmp_path / "analysis_outputs"

    generate_data(n=20, seed=123).to_csv(input_csv, index=False)
    analyse_survey(input_csv, analysis_dir)

    return analysis_dir


def test_report_contains_expected_sections_and_disclosure(tmp_path: Path):
    """The report should contain its main sections and synthetic-data warning."""
    analysis_dir = prepare_analysis(tmp_path)
    output_file = tmp_path / "research_report.docx"

    result = create_report(analysis_dir, output_file)

    assert result.exists()
    assert result.stat().st_size > 0

    document = Document(result)
    text = "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )

    for expected in [
        "Research and Data Skills Needs Assessment",
        "SYNTHETIC",
        "Executive Summary",
        "Methods",
        "Data Quality and Validation",
        "Descriptive Results",
        "Interpretation",
        "Limitations",
        "Recommendations for a Real Assessment",
        "Reproducibility and Provenance",
    ]:
        assert expected.lower() in text.lower(), (
            f"Expected report content not found: {expected}"
        )

    assert "20 synthetic records" in text
    assert (
        "not collected from actual educators or institutions"
        in text.lower()
    )


def test_report_includes_both_charts(tmp_path: Path):
    """Both analysis charts should be embedded in the Word document."""
    analysis_dir = prepare_analysis(tmp_path)
    output_file = tmp_path / "research_report.docx"

    create_report(analysis_dir, output_file)

    document = Document(output_file)
    assert len(document.inline_shapes) == 2


def test_report_refuses_incorrect_dataset_label(tmp_path: Path):
    """The report must refuse a manifest not labelled as synthetic."""
    analysis_dir = prepare_analysis(tmp_path)
    manifest_path = analysis_dir / "analysis_manifest.json"

    manifest = json.loads(
        manifest_path.read_text(encoding="utf-8")
    )
    manifest["dataset_type"] = "REAL"
    manifest_path.write_text(
        json.dumps(manifest, indent=2),
        encoding="utf-8",
    )

    output_file = tmp_path / "should_not_exist.docx"

    with pytest.raises(ValueError, match="SYNTHETIC_NOT_REAL"):
        create_report(analysis_dir, output_file)

    assert not output_file.exists()


def test_report_refuses_failed_validation(tmp_path: Path):
    """The report must not be generated when validation disallows analysis."""
    analysis_dir = prepare_analysis(tmp_path)
    validation_path = analysis_dir / "validation_report.json"

    validation = json.loads(
        validation_path.read_text(encoding="utf-8")
    )
    validation["analysis_allowed"] = False
    validation_path.write_text(
        json.dumps(validation, indent=2),
        encoding="utf-8",
    )

    output_file = tmp_path / "should_not_exist.docx"

    with pytest.raises(
        ValueError, match="validation does not allow analysis"
    ):
        create_report(analysis_dir, output_file)

    assert not output_file.exists()


def test_report_fails_if_required_table_is_missing(tmp_path: Path):
    """Missing analysis outputs should be reported rather than ignored."""
    analysis_dir = prepare_analysis(tmp_path)
    (analysis_dir / "q10_task_frequencies.csv").unlink()

    output_file = tmp_path / "incomplete_report.docx"

    with pytest.raises(
        FileNotFoundError, match="q10_task_frequencies"
    ):
        create_report(analysis_dir, output_file)

    assert not output_file.exists()
